"""Read-only Excel bibliography inspection and extraction."""

from __future__ import annotations

import math
import re
from pathlib import Path
from typing import Any
from urllib.parse import urlsplit
from zipfile import BadZipFile

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException


PROJECT_DIR = Path(__file__).resolve().parent
WORKSPACE_ROOT = PROJECT_DIR.parent
SUPPORTED_SUFFIXES = {".xlsx", ".xlsm"}
MAX_FILE_BYTES = 100 * 1024 * 1024
HEADER_SCAN_ROWS = 30
MAX_COLUMNS = 200
MAX_SHEET_SCAN_ROWS = 100_000
MAX_EXTRACT_ROWS = 5_000
DOI_PATTERN = re.compile(r"10\.\d{4,9}/[-._;()/:A-Z0-9]+", re.IGNORECASE)
OUTPUT_FIELDS = (
    "title", "authors", "first_author", "year", "journal/source", "DOI",
    "DOI_URL", "official_URL", "priority", "topic", "research_relationship",
)


def _error(code: str, message: str) -> dict[str, Any]:
    return {"ok": False, "source": "Excel", "error": {"code": code, "message": message}}


def _resolve_path(excel_path: str) -> tuple[Path | None, dict[str, Any] | None]:
    if not isinstance(excel_path, str) or not excel_path.strip():
        return None, _error("invalid_input", "excel_path 不能为空。")
    raw = Path(excel_path.strip())
    candidates = [raw] if raw.is_absolute() else [WORKSPACE_ROOT / raw, PROJECT_DIR / raw]
    selected = next((item for item in candidates if item.exists()), candidates[0])
    try:
        resolved = selected.resolve(strict=True)
        resolved.relative_to(WORKSPACE_ROOT.resolve())
    except (FileNotFoundError, OSError):
        return None, _error("file_not_found", "未找到指定的 Excel 文件。")
    except ValueError:
        return None, _error("path_outside_workspace", "只允许读取当前工作区内的 Excel 文件。")
    if not resolved.is_file():
        return None, _error("not_a_file", "excel_path 必须指向文件。")
    if resolved.suffix.lower() not in SUPPORTED_SUFFIXES:
        return None, _error("unsupported_format", "当前支持 .xlsx 和 .xlsm；旧版 .xls 需要先另存为 .xlsx。")
    try:
        if resolved.stat().st_size > MAX_FILE_BYTES:
            return None, _error("file_too_large", "Excel 文件超过 100 MiB 安全上限。")
    except OSError:
        return None, _error("file_unreadable", "无法读取 Excel 文件属性。")
    return resolved, None


def _open_workbook(excel_path: str):
    path, error = _resolve_path(excel_path)
    if error:
        return None, None, error
    try:
        workbook = load_workbook(
            filename=path, read_only=True, data_only=True, keep_links=False
        )
    except PermissionError:
        return None, None, _error("file_locked", "Excel 文件不可读；请检查文件权限或占用状态。")
    except (BadZipFile, InvalidFileException, OSError, ValueError):
        return None, None, _error("invalid_workbook", "文件不是可读取的有效 Excel 工作簿。")
    return workbook, path, None


def _cell_text(value: Any) -> str | None:
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, float):
        if not math.isfinite(value):
            return None
        if value.is_integer():
            return str(int(value))
    text = str(value).strip()
    return re.sub(r"\s+", " ", text) if text else None


def _normalise_header(value: Any) -> str:
    return re.sub(
        r"[\s_\-/\\|:：,，;；()（）\[\]【】]+", "", (_cell_text(value) or "").casefold()
    )


def _header_field(value: Any) -> str | None:
    header = _normalise_header(value)
    if not header:
        return None
    if any(x in header for x in ("第一作者", "首位作者", "firstauthor", "leadauthor")):
        return "first_author"
    if any(x in header for x in ("文献题目", "论文题目", "文章题目", "articletitle")) or header in {"标题", "题名", "title"}:
        return "title"
    if header in {"作者", "authors", "author", "全部作者", "作者列表"}:
        return "authors"
    if header in {"年份", "年", "year", "publicationyear", "发表年份", "出版年份"}:
        return "year"
    if any(x in header for x in ("期刊", "journal", "刊物", "来源期刊")) or header in {"source", "venue", "来源"}:
        return "journal/source"
    if "doi" in header and any(x in header for x in ("url", "链接", "网址", "完整")):
        return "mixed_url" if "官方" in header else "doi_url"
    if header == "doi" or header.startswith("doi号"):
        return "doi"
    if any(x in header for x in ("官方url", "官方完整url", "官方链接", "官方网址", "officialurl", "文献链接")) or header == "url":
        return "official_url"
    if any(x in header for x in ("优先级", "priority", "推荐层级")):
        return "priority"
    if any(x in header for x in ("与你研究的关系", "与研究的关系", "研究关系", "研究关联", "researchrelationship")):
        return "research_relationship"
    if any(x in header for x in ("主题", "topic", "关键词", "keyword")):
        return "topic"
    return None


def _header_mapping(row: tuple[Any, ...]) -> dict[int, str]:
    mapping: dict[int, str] = {}
    occupied: set[str] = set()
    for index, value in enumerate(row):
        field = _header_field(value)
        if field and (field == "mixed_url" or field not in occupied):
            mapping[index] = field
            occupied.add(field)
    return mapping


def _mapping_fields(mapping: dict[int, str]) -> set[str]:
    fields: set[str] = set()
    for field in mapping.values():
        if field == "doi":
            fields.add("DOI")
        elif field == "doi_url":
            fields.update({"DOI", "DOI_URL"})
        elif field == "official_url":
            fields.add("official_URL")
        elif field == "mixed_url":
            fields.update({"DOI", "DOI_URL", "official_URL"})
        else:
            fields.add(field)
    return fields


def _schema_score(mapping: dict[int, str]) -> int:
    weights = {
        "title": 6, "DOI": 5, "DOI_URL": 4, "authors": 3,
        "first_author": 3, "year": 2, "journal/source": 2,
        "official_URL": 1, "priority": 1, "topic": 1,
        "research_relationship": 1,
    }
    return sum(weights.get(field, 0) for field in _mapping_fields(mapping))


def _detect_header(worksheet) -> tuple[int | None, tuple[Any, ...], dict[int, str], int]:
    best: tuple[int | None, tuple[Any, ...], dict[int, str], int] = (None, (), {}, 0)
    for row_number, row in enumerate(
        worksheet.iter_rows(min_row=1, max_row=HEADER_SCAN_ROWS, max_col=MAX_COLUMNS, values_only=True), 1
    ):
        mapping = _header_mapping(row)
        score = _schema_score(mapping)
        if score > best[3]:
            best = (row_number, row, mapping, score)
    return best


def _is_reference_schema(mapping: dict[int, str]) -> bool:
    fields = _mapping_fields(mapping)
    return "title" in fields and bool(
        fields & {"DOI", "DOI_URL", "authors", "first_author", "year", "journal/source", "official_URL"}
    )


def _safe_filename(path: Path) -> str:
    return path.name


def list_sheets(excel_path: str) -> dict[str, Any]:
    workbook, path, error = _open_workbook(excel_path)
    if error:
        return error
    sheets: list[dict[str, Any]] = []
    try:
        for worksheet in workbook.worksheets:
            nonempty_rows = last_row = last_column = 0
            truncated = False
            for row_number, row in enumerate(worksheet.iter_rows(max_col=MAX_COLUMNS, values_only=True), 1):
                if row_number > MAX_SHEET_SCAN_ROWS:
                    truncated = True
                    break
                indexes = [i for i, value in enumerate(row, 1) if _cell_text(value) is not None]
                if indexes:
                    nonempty_rows += 1
                    last_row = row_number
                    last_column = max(last_column, max(indexes))
            header_row, _, mapping, score = _detect_header(worksheet)
            sheets.append({
                "name": worksheet.title,
                "nonempty_rows": nonempty_rows,
                "last_nonempty_row": last_row,
                "last_nonempty_column": last_column,
                "scan_truncated": truncated,
                "detected_header_row": header_row,
                "reference_candidate": _is_reference_schema(mapping),
                "reference_score": score,
            })
    finally:
        workbook.close()
    return {"ok": True, "source": "Excel", "file": _safe_filename(path), "count": len(sheets), "sheets": sheets, "read_only": True}


def _schema_for_sheet(worksheet) -> dict[str, Any]:
    header_row, headers, mapping, score = _detect_header(worksheet)
    columns = []
    for index, field in mapping.items():
        columns.append({
            "column": index + 1,
            "header": _cell_text(headers[index]) if index < len(headers) else None,
            "mapped_fields": (
                ["DOI", "DOI_URL", "official_URL"]
                if field == "mixed_url"
                else ["DOI", "DOI_URL"]
                if field == "doi_url"
                else ["DOI"]
                if field == "doi"
                else ["official_URL"]
                if field == "official_url"
                else [field]
            ),
        })
    fields = _mapping_fields(mapping)
    return {
        "sheet": worksheet.title,
        "header_row": header_row,
        "recognized_fields": sorted(fields),
        "columns": columns,
        "reference_candidate": _is_reference_schema(mapping),
        "reference_score": score,
        "missing_recommended_fields": [field for field in OUTPUT_FIELDS if field not in fields],
    }


def inspect_schema(excel_path: str, sheet_name: str | None = None) -> dict[str, Any]:
    workbook, path, error = _open_workbook(excel_path)
    if error:
        return error
    try:
        if sheet_name is not None and sheet_name not in workbook.sheetnames:
            return _error("sheet_not_found", "未找到指定的 Sheet。")
        worksheets = [workbook[sheet_name]] if sheet_name else list(workbook.worksheets)
        schemas = [_schema_for_sheet(item) for item in worksheets]
    finally:
        workbook.close()
    candidates = [item for item in schemas if item["reference_candidate"]]
    recommended = max(candidates, key=lambda item: item["reference_score"], default=None)
    return {
        "ok": True, "source": "Excel", "file": _safe_filename(path),
        "schemas": schemas, "recommended_sheet": recommended["sheet"] if recommended else None,
        "read_only": True,
    }


def _extract_doi(value: Any) -> str | None:
    match = DOI_PATTERN.search(_cell_text(value) or "")
    return match.group(0).rstrip(".,;:)]}〉》。，；：").lower() if match else None


def _valid_http_url(value: Any) -> str | None:
    text = _cell_text(value)
    if not text:
        return None
    try:
        parsed = urlsplit(text)
    except ValueError:
        return None
    return text if parsed.scheme.lower() in {"http", "https"} and parsed.hostname else None


def _authors(value: Any) -> list[str]:
    text = _cell_text(value)
    if not text:
        return []
    if any(mark in text for mark in (";", "；", "|", "、")):
        return [part.strip() for part in re.split(r"[;；|、]", text) if part.strip()]
    return [text]


def _year(value: Any) -> tuple[int | None, str | None]:
    if value in (None, ""):
        return None, None
    if isinstance(value, int) and not isinstance(value, bool) and 1000 <= value <= 3000:
        return value, None
    if isinstance(value, float) and value.is_integer() and 1000 <= value <= 3000:
        return int(value), None
    match = re.search(r"\b(?:1|2)\d{3}\b", _cell_text(value) or "")
    return (int(match.group(0)), None) if match else (None, "年份无法识别")


def _priority(value: Any) -> tuple[int | float | str | None, str | None]:
    if value in (None, ""):
        return None, None
    if isinstance(value, bool):
        return None, "优先级无法识别"
    if isinstance(value, (int, float)) and math.isfinite(float(value)):
        return value, None
    text = _cell_text(value)
    try:
        number = float(text)
        return (int(number) if number.is_integer() else number), None
    except (TypeError, ValueError):
        return text, "优先级不是数字，已保留原值"


def _raw_row(headers: tuple[Any, ...], row: tuple[Any, ...]) -> dict[str, str]:
    result: dict[str, str] = {}
    for index, value in enumerate(row):
        text = _cell_text(value)
        if text is None:
            continue
        key = _cell_text(headers[index]) if index < len(headers) else None
        key = key or f"column_{index + 1}"
        if key in result:
            key = f"{key}_{index + 1}"
        result[key] = text
    return result


def _empty_record(sheet: str, source_row: int) -> dict[str, Any]:
    return {
        "source_sheet": sheet, "source_row": source_row,
        "title": None, "authors": [], "first_author": None, "year": None,
        "journal/source": None, "DOI": None, "DOI_URL": None,
        "official_URL": None, "priority": None, "topic": None,
        "research_relationship": None, "errors": [],
    }


def _parse_row(sheet: str, source_row: int, row: tuple[Any, ...], mapping: dict[int, str]) -> dict[str, Any]:
    record = _empty_record(sheet, source_row)
    for index, field in mapping.items():
        value = row[index] if index < len(row) else None
        if field == "title":
            record["title"] = _cell_text(value)
        elif field == "authors":
            record["authors"] = _authors(value)
        elif field == "first_author":
            record["first_author"] = _cell_text(value)
        elif field == "year":
            record["year"], warning = _year(value)
            if warning:
                record["errors"].append(warning)
        elif field == "journal/source":
            record["journal/source"] = _cell_text(value)
        elif field == "doi":
            record["DOI"] = _extract_doi(value)
            if _cell_text(value) and not record["DOI"]:
                record["errors"].append("DOI 无法识别")
        elif field in {"doi_url", "mixed_url"}:
            text, doi = _cell_text(value), _extract_doi(value)
            if doi:
                record["DOI"] = record["DOI"] or doi
                record["DOI_URL"] = f"https://doi.org/{doi}"
            elif text and field == "mixed_url" and _valid_http_url(text):
                record["official_URL"] = text
            elif text:
                record["errors"].append("DOI URL 无法识别")
        elif field == "official_url":
            text, doi = _cell_text(value), _extract_doi(value)
            if doi:
                record["DOI"] = record["DOI"] or doi
                record["DOI_URL"] = f"https://doi.org/{doi}"
            elif text:
                record["official_URL"] = _valid_http_url(text)
                if not record["official_URL"]:
                    record["errors"].append("官方 URL 无法识别")
        elif field == "priority":
            record["priority"], warning = _priority(value)
            if warning:
                record["errors"].append(warning)
        elif field == "topic":
            record["topic"] = _cell_text(value)
        elif field == "research_relationship":
            record["research_relationship"] = _cell_text(value)
    if record["DOI"] and not record["DOI_URL"]:
        record["DOI_URL"] = f"https://doi.org/{record['DOI']}"
    if not record["first_author"] and record["authors"]:
        record["first_author"] = record["authors"][0]
    return record


def _sheet_extraction(worksheet, max_rows: int) -> dict[str, Any]:
    header_row, headers, mapping, score = _detect_header(worksheet)
    if header_row is None or not _is_reference_schema(mapping):
        return {
            "sheet": worksheet.title, "header_row": header_row,
            "reference_candidate": False, "reference_score": score,
            "records": [], "unrecognized_records": [],
            "errors": ["未识别到可用的文献记录表头"],
            "scanned_data_rows": 0, "truncated": False,
        }
    records, unrecognized = [], []
    scanned = 0
    for source_row, row in enumerate(
        worksheet.iter_rows(
            min_row=header_row + 1, max_row=header_row + max_rows + 1,
            max_col=max(mapping) + 1, values_only=True,
        ), header_row + 1,
    ):
        if not any(_cell_text(value) is not None for value in row):
            continue
        scanned += 1
        if scanned > max_rows:
            break
        try:
            record = _parse_row(worksheet.title, source_row, row, mapping)
        except Exception:
            unrecognized.append({
                "source_sheet": worksheet.title, "source_row": source_row,
                "reason": "该行发生未预期的解析错误", "raw_values": _raw_row(headers, row),
            })
            continue
        if record["title"] or record["DOI"]:
            records.append(record)
        else:
            unrecognized.append({
                "source_sheet": worksheet.title, "source_row": source_row,
                "reason": "缺少可识别的标题和 DOI", "raw_values": _raw_row(headers, row),
            })
    return {
        "sheet": worksheet.title, "header_row": header_row,
        "reference_candidate": True, "reference_score": score,
        "records": records, "unrecognized_records": unrecognized, "errors": [],
        "scanned_data_rows": scanned, "truncated": scanned > max_rows,
    }


def _best_reference_sheet(workbook):
    candidates = []
    for worksheet in workbook.worksheets:
        _, _, mapping, score = _detect_header(worksheet)
        if _is_reference_schema(mapping):
            candidates.append((score, worksheet))
    return max(candidates, key=lambda item: item[0])[1] if candidates else None


def extract(excel_path: str, sheet_name: str | None = None, max_rows: int = 2_000) -> dict[str, Any]:
    if not isinstance(max_rows, int) or isinstance(max_rows, bool) or max_rows < 1:
        return _error("invalid_input", "max_rows 必须是正整数。")
    max_rows = min(max_rows, MAX_EXTRACT_ROWS)
    workbook, path, error = _open_workbook(excel_path)
    if error:
        return error
    try:
        if sheet_name is None:
            worksheet = _best_reference_sheet(workbook)
            if worksheet is None:
                return _error("no_reference_sheet", "没有发现可识别的文献目录 Sheet。")
        elif sheet_name in workbook.sheetnames:
            worksheet = workbook[sheet_name]
        else:
            return _error("sheet_not_found", "未找到指定的 Sheet。")
        result = _sheet_extraction(worksheet, max_rows)
    finally:
        workbook.close()
    records = result.pop("records")
    unrecognized = result.pop("unrecognized_records")
    return {
        "ok": True, "source": "Excel", "file": _safe_filename(path), **result,
        "record_count": len(records),
        "records_with_errors": sum(bool(item["errors"]) for item in records),
        "unrecognized_count": len(unrecognized),
        "records": records, "unrecognized_records": unrecognized, "read_only": True,
    }


def preview(excel_path: str, sheet_name: str | None = None, limit: int = 10) -> dict[str, Any]:
    if not isinstance(limit, int) or isinstance(limit, bool) or limit < 1:
        return _error("invalid_input", "limit 必须是正整数。")
    limit = min(limit, 50)
    result = extract(excel_path, sheet_name, max_rows=MAX_EXTRACT_ROWS)
    if not result.get("ok"):
        return result
    records = result.pop("records")
    unrecognized = result.pop("unrecognized_records")
    return {
        **result, "preview_count": min(limit, len(records)),
        "records": records[:limit], "unrecognized_preview": unrecognized[:limit],
    }
